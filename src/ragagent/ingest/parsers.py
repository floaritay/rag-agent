"""六格式解析器：md/txt/docx/pptx/xlsx/pdf → Block 列表。

约定：
- .doc/.ppt 不做自动转换，抛 UnsupportedFormatError 并提示手动另存；
- 表格统一转为 Markdown 文本（type="table"），解析失败的 PDF 表格退化为页文本并记 warning；
- 图片：pdf/docx/pptx 抽出为 ImageItem（字节+位置+alt），由 images.py 落盘与
  文字化（IMAGE_DESCRIBE 开关）；md/txt/xlsx 不抽图。image_count 仍为发现总数。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path


class UnsupportedFormatError(Exception):
    """不支持的格式（含 doc/ppt 需手动另存的引导提示）。"""


@dataclass
class Block:
    type: str  # text / table / code
    text: str
    heading_path: list[str] = field(default_factory=list)
    page: int | None = None
    line_start: int | None = None
    line_end: int | None = None
    image_path: str | None = None  # caption 块携带的图片（相对 storage 根）


@dataclass
class ImageItem:
    """解析器抽出的原始图片（尚未落盘）。"""

    bytes: bytes
    ext: str                     # "png"/"jpeg"…（无点）
    page: int | None             # PDF 页号 / pptx slide 序号 / docx None
    heading_path: list[str]
    ordinal: int                 # 文档内 0-based（落盘文件名）
    alt: str = ""                # docx/pptx 图片替代文字
    insert_at: int = 0           # caption 块插入点（发现时 len(doc.blocks)）


@dataclass
class ParsedDoc:
    blocks: list[Block] = field(default_factory=list)
    image_count: int = 0
    warnings: list[str] = field(default_factory=list)
    images: list[ImageItem] = field(default_factory=list)


SUPPORTED_EXTS = {".md", ".markdown", ".txt", ".pdf", ".docx", ".pptx", ".xlsx"}


def parse(path: Path) -> ParsedDoc:
    ext = path.suffix.lower()
    if ext in (".doc", ".ppt"):
        raise UnsupportedFormatError(
            "暂不支持 .doc/.ppt：请手动另存为 .docx/.pptx 后重新入库"
        )
    if ext in (".md", ".markdown"):
        return parse_markdown(path)
    if ext == ".txt":
        return parse_text(path)
    if ext == ".pdf":
        return parse_pdf(path)
    if ext == ".docx":
        return parse_docx(path)
    if ext == ".pptx":
        return parse_pptx(path)
    if ext == ".xlsx":
        return parse_xlsx(path)
    raise UnsupportedFormatError(f"不支持的格式 {ext}（支持：md/txt/docx/pptx/xlsx/pdf）")


# ---------- 表格 → Markdown ----------

def _cell(v) -> str:
    if v is None:
        return ""
    return str(v).replace("\r", " ").replace("\n", " ").replace("|", "\\|").strip()


def rows_to_markdown(header: list, rows: list[list]) -> str:
    header = [_cell(h) for h in header]
    if not header:
        header = [""]
    width = len(header)
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    for r in rows:
        cells = [_cell(c) for c in r]
        cells = (cells + [""] * width)[:width]
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


# ---------- markdown / txt ----------

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")


def _blocks_from_lines(lines: list[str]) -> list[Block]:
    """markdown 行流 → Block 列表（parse_markdown 与 MinerU 通道共用的核心）。"""
    blocks: list[Block] = []
    stack: list[tuple[int, str]] = []

    def hp() -> list[str]:
        return [t for _, t in stack]

    i = 0
    para: list[str] = []
    para_start = 1

    def flush_para(end_line: int) -> None:
        nonlocal para
        if para:
            text = "\n".join(para).strip()
            if text:
                blocks.append(
                    Block("text", text, hp(), None, para_start, end_line)
                )
        para = []

    while i < len(lines):
        line = lines[i]
        m = _HEADING_RE.match(line)
        if m:
            flush_para(max(i, 1))
            level, title = len(m.group(1)), m.group(2).strip()
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, title))
            i += 1
            continue
        if line.strip().startswith("```"):
            flush_para(max(i, 1))
            start = i + 1
            i += 1
            buf: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1  # 跳过收尾 ```
            blocks.append(
                Block("code", "\n".join(buf), hp(), None, start, i)
            )
            continue
        if line.lstrip().startswith("|"):
            flush_para(max(i, 1))
            start = i + 1
            buf = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                buf.append(lines[i].strip())
                i += 1
            rows = [
                [c.strip() for c in row.strip().strip("|").split("|")]
                for row in buf
            ]
            # 去掉形如 | --- | --- | 的分隔行
            rows = [
                r for r in rows
                if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in r)
            ]
            if rows:
                blocks.append(
                    Block("table", rows_to_markdown(rows[0], rows[1:]), hp(), None, start, i)
                )
            continue
        if line.strip():
            if not para:
                para_start = i + 1
            para.append(line)
        else:
            flush_para(max(i, 1))
        i += 1
    flush_para(len(lines))
    return blocks


def parse_markdown(path: Path) -> ParsedDoc:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    doc = ParsedDoc()
    doc.blocks = _blocks_from_lines(lines)
    return doc


def parse_text(path: Path) -> ParsedDoc:
    doc = ParsedDoc()
    buf: list[str] = []
    start = 1
    for lineno, line in enumerate(
        path.read_text(encoding="utf-8", errors="replace").splitlines(), 1
    ):
        if line.strip():
            if not buf:
                start = lineno
            buf.append(line)
        elif buf:
            doc.blocks.append(Block("text", "\n".join(buf), [], None, start, lineno - 1))
            buf = []
    if buf:
        doc.blocks.append(Block("text", "\n".join(buf), [], None, start, None))
    return doc


# ---------- pdf ----------

def parse_pdf(path: Path) -> ParsedDoc:
    import pymupdf

    doc = ParsedDoc()
    with pymupdf.open(str(path)) as pdf:
        seen_xrefs: set[int] = set()  # 同一图片对象跨页复用只抽一次
        for pno, page in enumerate(pdf, 1):
            xrefs = page.get_images(full=True)
            doc.image_count += len(xrefs)
            for img in xrefs:
                xref = img[0]
                if xref in seen_xrefs:
                    continue
                seen_xrefs.add(xref)
                try:
                    # PyMuPDF 1.28：字节在 "image" 键（不是 "bytes"）
                    info = pdf.extract_image(xref)
                except Exception:  # noqa: BLE001 —— 单图失败不影响解析
                    continue
                data = info.get("image")
                if not data:
                    continue
                doc.images.append(ImageItem(
                    bytes=data,
                    ext=(info.get("ext") or "png").lower(),
                    page=pno,
                    heading_path=[],
                    ordinal=len(doc.images),
                    insert_at=len(doc.blocks),
                ))
            table_rects: list = []
            try:
                tf = page.find_tables()
                for t in tf.tables:
                    data = t.extract() or []
                    data = [r for r in data if any(c for c in r)]
                    if not data:
                        continue
                    doc.blocks.append(
                        Block("table", rows_to_markdown(data[0], data[1:]), [], pno)
                    )
                    table_rects.append(pymupdf.Rect(t.bbox))
            except Exception as e:  # noqa: BLE001 —— 表格提取失败退化为页文本
                doc.warnings.append(f"第{pno}页表格提取失败，已退化为页文本：{e}")
            for b in page.get_text("blocks"):
                if b[6] != 0:  # 非文本块（图片）
                    continue
                rect = pymupdf.Rect(b[0], b[1], b[2], b[3])
                if any(rect.intersects(tr) for tr in table_rects):
                    continue
                text = b[4].strip()
                if text:
                    doc.blocks.append(Block("text", text, [], pno))
    return doc


# ---------- docx ----------

def parse_docx(path: Path) -> ParsedDoc:
    import docx as docx_mod
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx_mod.Document(str(path))
    doc = ParsedDoc()
    doc.image_count = len(d.inline_shapes)
    stack: list[tuple[int, str]] = []

    def hp() -> list[str]:
        return [t for _, t in stack]

    def add_docx_images(element) -> None:
        """段落/表格里的内联与浮动图（a:blip 锚点）→ ImageItem（含 descr alt）。"""
        for r_id in element.xpath(".//a:blip/@r:embed"):
            try:
                part = d.part.rels[r_id].target_part
            except KeyError:
                continue
            data = part.blob
            if not data:
                continue
            ctype = getattr(part, "content_type", "") or ""
            ext = ctype.split("/", 1)[1].lower() if "/" in ctype else "png"
            if ext == "jpg":
                ext = "jpeg"
            descrs = element.xpath(".//wp:docPr/@descr")
            alt = (descrs[0] or "").strip() if descrs else ""
            doc.images.append(ImageItem(
                bytes=data,
                ext=ext,
                page=None,
                heading_path=hp(),
                ordinal=len(doc.images),
                alt=alt,
                insert_at=len(doc.blocks),
            ))

    for child in d.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, d)
            text = para.text.strip()
            style = (para.style.name if para.style is not None else "") or ""
            m = re.match(r"^(?:Heading|标题)\s*(\d)", style)
            if m and text:
                level = int(m.group(1))
                while stack and stack[-1][0] >= level:
                    stack.pop()
                stack.append((level, text))
            elif text:
                doc.blocks.append(Block("text", text, hp()))
            add_docx_images(child)
        elif child.tag == qn("w:tbl"):
            tbl = Table(child, d)
            rows = [[cell.text for cell in row.cells] for row in tbl.rows]
            if rows:
                doc.blocks.append(
                    Block("table", rows_to_markdown(rows[0], rows[1:]), hp())
                )
            add_docx_images(child)
    return doc


# ---------- pptx ----------

def parse_pptx(path: Path) -> ParsedDoc:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    doc = ParsedDoc()
    for idx, slide in enumerate(prs.slides, 1):
        slide_hp = [f"第{idx}页"]
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                doc.image_count += 1
                try:
                    img = shape.image
                    data = img.blob
                    ext = (img.ext or "png").lower()
                except Exception:  # noqa: BLE001 —— 单图失败不影响解析
                    continue
                if not data:
                    continue
                try:
                    alt = (shape._element.nvPicPr.cNvPr.get("descr") or "").strip()
                except Exception:  # noqa: BLE001
                    alt = ""
                doc.images.append(ImageItem(
                    bytes=data,
                    ext=ext,
                    page=idx,
                    heading_path=list(slide_hp),
                    ordinal=len(doc.images),
                    alt=alt,
                    insert_at=len(doc.blocks),
                ))
                continue
            if getattr(shape, "has_table", False):
                tbl = shape.table
                rows = [[_cell(c.text) for c in r.cells] for r in tbl.rows]
                if rows:
                    doc.blocks.append(
                        Block("table", rows_to_markdown(rows[0], rows[1:]), slide_hp, idx)
                    )
                continue
            if shape.has_text_frame:
                txt = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
                if txt:
                    doc.blocks.append(Block("text", txt, slide_hp, idx))
    return doc


# ---------- xlsx ----------

_XLSX_ROW_BATCH = 40  # 每 40 行一块，每块重复表头


def parse_xlsx(path: Path) -> ParsedDoc:
    from openpyxl import load_workbook

    doc = ParsedDoc()
    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = [
                [_cell(c) for c in row]
                for row in ws.iter_rows(values_only=True)
            ]
            rows = [r for r in rows if any(c for c in r)]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            header, data = rows[0], rows[1:]
            if not data:
                doc.blocks.append(
                    Block("table", rows_to_markdown(header, []), [ws.title])
                )
            for i in range(0, len(data), _XLSX_ROW_BATCH):
                doc.blocks.append(
                    Block(
                        "table",
                        rows_to_markdown(header, data[i : i + _XLSX_ROW_BATCH]),
                        [ws.title],
                    )
                )
    finally:
        wb.close()
    return doc
